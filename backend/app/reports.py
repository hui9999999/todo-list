"""Filtered task reports. Full restoration backups remain in data.py."""
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from html import escape
from io import BytesIO, StringIO
from math import ceil
from typing import Literal
from zoneinfo import ZoneInfo
import csv
import re

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from sqlalchemy import select, func, or_, and_
from sqlalchemy.orm import Session
from .auth import current_user
from .db import database
from .models import Task, User, utcnow
from .advanced_models import Project

router = APIRouter()
SCOPES = {'work': '工作', 'life': '生活', 'study': '学习', 'unassigned': '未分类'}
STATES = {'todo': '待办', 'doing': '进行中', 'done': '已完成', 'failed': '失败', 'cancelled': '已取消'}
PRIORITIES = {'high': '高', 'mid': '中', 'low': '低'}
BASES = {'created': '创建时间', 'completed': '完成时间', 'updated': '更新时间', 'due': '截止时间'}


@dataclass
class Filters:
    start: date | None
    end: date | None
    date_field: str
    scope: str
    status: str
    include_trash: bool
    include_notes: bool


def report_filters(start: date | None = None, end: date | None = None,
                   date_field: Literal['created', 'completed', 'updated', 'due'] = 'created',
                   scope: Literal['all', 'work', 'life', 'study', 'unassigned'] = 'all',
                   status: Literal['all', 'todo', 'doing', 'done', 'failed', 'cancelled'] = 'all',
                   include_trash: bool = False, include_notes: bool = True):
    if start and end and start > end:
        raise HTTPException(422, '结束日期不能早于开始日期')
    if end == date.max:
        raise HTTPException(422, '结束日期超出可导出范围')
    return Filters(start, end, date_field, scope, status, include_trash, include_notes)


def select_report(db: Session, user: User, filters: Filters):
    zone = ZoneInfo(user.timezone)
    query = select(Task).where(Task.user_id == user.id)
    if not filters.include_trash:
        query = query.where(Task.deleted_at.is_(None))
    if filters.scope != 'all':
        query = query.where(Task.scope == filters.scope)
    if filters.status != 'all':
        query = query.where(Task.status == filters.status)
    dates = {'created': func.coalesce(Task.legacy_created_at, Task.created_at),
             'completed': Task.completed_at, 'updated': Task.updated_at, 'due': Task.due_at}
    column = dates[filters.date_field]
    if filters.date_field in ('completed', 'due'):
        query = query.where(or_(Task.due_at.is_not(None), Task.due_date.is_not(None))
                            if filters.date_field == 'due' else column.is_not(None))
    for boundary, lower in ((filters.start, True), (filters.end, False)):
        if not boundary:
            continue
        local_date = boundary if lower else boundary + timedelta(days=1)
        instant = datetime.combine(local_date, time.min, zone)
        condition = column >= instant if lower else column < instant
        if filters.date_field == 'due':
            date_condition = Task.due_date >= local_date if lower else Task.due_date < local_date
            condition = or_(condition, and_(Task.due_at.is_(None), date_condition))
        query = query.where(condition)
    tasks = list(db.scalars(query.order_by(func.coalesce(Task.legacy_created_at, Task.created_at), Task.id).limit(10001)))
    if len(tasks) > 10000:
        raise HTTPException(422, '单次最多导出 10000 条任务，请缩小时间范围或选择分类')
    projects = {p.id: p.title for p in db.scalars(select(Project).where(Project.user_id == user.id))}
    def stamp(value):
        return value.astimezone(zone).replace(tzinfo=None) if value else None
    rows = [{
        'id': t.id, 'title': t.title, 'scope': SCOPES[t.scope], 'status': STATES[t.status],
        'priority': PRIORITIES[t.priority], 'project': projects.get(t.project_id, ''),
        'parent_id': t.parent_id or '', 'created_at': stamp(t.legacy_created_at or t.created_at),
        'updated_at': stamp(t.updated_at), 'completed_at': stamp(t.completed_at),
        'due_date': t.due_date, 'due_at': stamp(t.due_at), 'estimated_minutes': t.estimated_minutes,
        'notes': t.notes if filters.include_notes else '', 'fail_reason': t.fail_reason,
        'hidden': '是' if t.hidden else '否', 'trash': '是' if t.deleted_at else '否',
    } for t in tasks]
    summary = {'total': len(tasks), 'scopes': {key: sum(t.scope == key for t in tasks) for key in SCOPES},
               'states': {key: sum(t.status == key for t in tasks) for key in STATES}}
    description = f"按{BASES[filters.date_field]}筛选：{filters.start or '不限'} 至 {filters.end or '不限'}（含结束当天）"
    selection = f"分类：{SCOPES.get(filters.scope, '全部')}；状态：{STATES.get(filters.status, '全部')}；垃圾站：{'包含' if filters.include_trash else '不包含'}"
    return {'rows': rows, 'summary': summary, 'description': description, 'selection': selection,
            'timezone': user.timezone, 'generated_at': utcnow().astimezone(zone).strftime('%Y-%m-%d %H:%M'),
            'include_notes': filters.include_notes}


@router.get('/exports/preview')
def preview_report(filters: Filters = Depends(report_filters), user: User = Depends(current_user),
                   db: Session = Depends(database, scope='function')):
    report = select_report(db, user, filters)
    return {key: value for key, value in report.items() if key != 'rows'}


HEADERS = [('title', '任务名称'), ('scope', '分类'), ('status', '状态'), ('priority', '优先级'),
           ('project', '项目'), ('created_at', '创建时间'), ('updated_at', '更新时间'),
           ('completed_at', '完成时间'), ('due_date', '截止日期'), ('due_at', '截止时刻'),
           ('estimated_minutes', '预计分钟'), ('notes', '备注'), ('fail_reason', '失败原因'),
           ('hidden', '已隐藏'), ('trash', '垃圾站'), ('id', '任务 ID'), ('parent_id', '父任务 ID')]


def display(value):
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(value, date):
        return value.isoformat()
    return '' if value is None else str(value)


def clean(value):
    # XML excludes control characters even when the underlying task allows them.
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '\ufffd', str(value))


def excel_report(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, Reference
    from openpyxl.utils import get_column_letter
    book = Workbook()
    summary = book.active
    summary.title = '导出概览'
    summary.append([report.get('title','任务数据导出')])
    summary.append([report['description']])
    summary.append([report['selection']])
    summary.append([f"时区：{report['timezone']}；导出于 {report['generated_at']}"])
    summary.append(['任务总数', report['summary']['total']])
    if report.get('period_summary'):
        period=book.create_sheet('周期统计');period.append([report['title']])
        for key,label in (('completed_tasks','完成任务'),('completion_events','完成事件'),('created_tasks','新增任务'),('focus_seconds','专注秒数')):
            period.append([label,report['period_summary'][key]])
        daily=book.create_sheet('每日统计');daily.append(['日期','新增任务','完成事件','专注秒数'])
        for day in report['daily']:daily.append([day['date'],day['created'],day['completed'],day['focus_seconds']])
    summary.append([])
    summary.append(['分类', '任务数', None, '状态', '任务数'])
    for i in range(5):
        scope = list(SCOPES)[i] if i < len(SCOPES) else None
        state = list(STATES)[i]
        summary.append([SCOPES[scope] if scope else None, report['summary']['scopes'][scope] if scope else None,
                        None, STATES[state], report['summary']['states'][state]])
    for col in ('A', 'D'):
        summary.column_dimensions[col].width = 22
    for col in ('B', 'C', 'E'):
        summary.column_dimensions[col].width = 14
    for row in range(1, 5):
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        summary.cell(row, 1).alignment = Alignment(wrap_text=True, vertical='center')
        summary.row_dimensions[row].height = 32
    summary['A1'].font = Font(size=20, bold=True, color='263D33')
    for title, column, last_row, anchor in [('分类分布', 1, 11, 'A15'), ('状态分布', 4, 12, 'J15')]:
        chart = BarChart()
        chart.type = 'bar'
        chart.title = title
        chart.y_axis.title = '任务数'
        chart.add_data(Reference(summary, min_col=column + 1, min_row=7, max_row=last_row), titles_from_data=True)
        chart.set_categories(Reference(summary, min_col=column, min_row=8, max_row=last_row))
        chart.legend = None
        chart.width, chart.height = 15, 8
        counts = list(report['summary']['scopes' if column == 1 else 'states'].values())
        largest = max(counts, default=0)
        unit = max(1, ceil(largest / 5))
        chart.y_axis.numFmt = '0'
        chart.y_axis.majorUnit = unit
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = max(unit, ceil(largest / unit) * unit)
        chart.series[0].graphicalProperties.solidFill = '48775F'
        chart.series[0].graphicalProperties.line.solidFill = '48775F'
        summary.add_chart(chart, anchor)
    sheet = book.create_sheet('任务清单')
    sheet.append([label for _, label in HEADERS])
    for row in report['rows']:
        sheet.append([clean(row[key]) if isinstance(row[key], str) else row[key] for key, _ in HEADERS])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Preserve literal user text, including leading = + - @.
            if isinstance(cell.value, datetime):
                cell.number_format = 'yyyy-mm-dd hh:mm:ss'
            elif isinstance(cell.value, date):
                cell.number_format = 'yyyy-mm-dd'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    widths = [40, 12, 14, 12, 24, 24, 24, 24, 16, 24, 14, 65, 36, 12, 12, 38, 38]
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = 'B2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = '1:1'
    for target, row in ((sheet, 1), (summary, 7)):
        for cell in target[row]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='374A43')
        target.row_dimensions[row].height = 26
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def chart_png(report, kind='scopes'):
    """Native CJK fonts supplied by MuPDF keep charts portable on Windows and Linux."""
    import pymupdf as pdf
    doc = pdf.open()
    page = doc.new_page(width=500, height=260)
    maximum = max(report['summary']['total'], 1)
    for i, (key, label) in enumerate((SCOPES if kind == 'scopes' else STATES).items()):
        y = 44 + i * 48
        count = report['summary'][kind][key]
        page.insert_text((12, y), label, fontname='china-s', fontsize=12)
        page.draw_rect(pdf.Rect(76, y - 12, 428, y + 2), fill=(.92, .94, .93), color=None)
        if count:
            page.draw_rect(pdf.Rect(76, y - 12, 76 + 352 * count / maximum, y + 2), fill=(.28, .46, .37), color=None)
        page.insert_text((444, y), str(count), fontsize=12)
    content = page.get_pixmap(matrix=pdf.Matrix(2, 2)).tobytes('png')
    doc.close()
    return content


def details(row, report):
    fields = ['分类', '状态', '优先级', '项目', '创建时间', '更新时间', '完成时间', '截止日期', '截止时刻',
              '预计分钟', '失败原因', '已隐藏', '垃圾站', '任务 ID', '父任务 ID']
    values = {label: display(row[key]) for key, label in HEADERS}
    return [(label, values[label]) for label in fields if values[label]]


def word_report(report):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.85)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].paragraph_format.space_after = Pt(7)
    doc.add_paragraph(report.get('title','任务数据导出'), 'Title')
    doc.add_paragraph(report['description'])
    doc.add_paragraph(report['selection'])
    doc.add_paragraph(f"时区：{report['timezone']}；导出于 {report['generated_at']}")
    doc.add_paragraph(f"共 {report['summary']['total']} 条任务。" + '；'.join(f"{STATES[k]} {v} 条" for k, v in report['summary']['states'].items()) + '。')
    doc.add_heading('分类分布', 1)
    doc.add_picture(BytesIO(chart_png(report)), width=Inches(6.1))
    doc.add_heading('状态分布', 1)
    doc.add_picture(BytesIO(chart_png(report, 'states')), width=Inches(6.1))
    doc.add_heading('任务明细', 1)
    if not report['rows']:
        doc.add_paragraph('没有符合筛选条件的任务。')
    for i, row in enumerate(report['rows'], 1):
        doc.add_heading(clean(f"{i} {row['title']}"), 2)
        # Short paragraphs flow naturally across pages; long notes do not squeeze into a table cell.
        for label, value in details(row, report):
            doc.add_paragraph(clean(f'{label}：{value}'))
        if report['include_notes'] and row['notes']:
            doc.add_paragraph('备注')
            for paragraph in row['notes'].split('\n'):
                doc.add_paragraph(clean(paragraph))
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('第 ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    footer.add_run(' 页')
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def pdf_report(report):
    import pymupdf as pdf
    def safe(value):
        # MuPDF does not wrap oversized unbroken tokens with CSS overflow-wrap.
        # Invisible break opportunities preserve every visible character in long URLs/text.
        value = re.sub(r'\S{41,}', lambda match: '\u200b'.join(
            match[0][offset:offset + 40] for offset in range(0, len(match[0]), 40)), clean(value))
        return escape(value)
    blocks = [f"<h1>{safe(report.get('title','任务数据导出'))}</h1>", f"<p>{safe(report['description'])}</p>", f"<p>{safe(report['selection'])}</p>",
              f"<p>时区：{safe(report['timezone'])}；导出于 {report['generated_at']}</p>",
              f"<p>共 {report['summary']['total']} 条任务。" + '；'.join(f"{STATES[k]} {v} 条" for k, v in report['summary']['states'].items()) + '。</p>',
              '<h2>分类分布</h2>']
    for kind, labels in (('scopes', SCOPES), ('states', STATES)):
        if kind == 'states':
            blocks.append('<h2>状态分布</h2>')
        for key, label in labels.items():
            count = report['summary'][kind][key]
            percentage = count / max(report['summary']['total'], 1) * 100
            blocks.append(f'<p>{label} {count} 条</p><div class="track"><div style="width:{percentage}%;background:#48775f;height:9pt"></div></div>')
    blocks.append('<h2>任务明细</h2>')
    if not report['rows']:
        blocks.append('<p>没有符合筛选条件的任务。</p>')
    for i, row in enumerate(report['rows'], 1):
        blocks.append(f"<h3>{i} {safe(row['title'])}</h3>")
        blocks.append('<p>' + '<br>'.join(safe(f'{label}：{value}') for label, value in details(row, report)) + '</p>')
        if report['include_notes'] and row['notes']:
            blocks.append('<p>备注<br>' + '<br>'.join(safe(line) for line in row['notes'].split('\n')) + '</p>')
    story = pdf.Story(''.join(blocks), user_css='body{font-family:sans-serif;font-size:10pt;line-height:1.55;color:#24372e}h1{font-size:24pt;color:#000}h2{font-size:16pt;color:#000;margin-top:18pt}h3{font-size:12pt;color:#000;page-break-after:avoid}p{margin:6pt 0;overflow-wrap:anywhere}.track{background:#edf1ee;height:9pt;width:100%}')
    box = pdf.Rect(0, 0, 595, 842)
    doc = story.write_with_links(lambda number, filled: (box, pdf.Rect(50, 48, 545, 790), None))
    for index, page in enumerate(doc):
        page.insert_text((490, 815), f'{index + 1} / {len(doc)}', fontsize=9, color=(.4, .45, .42))
    result = doc.tobytes(garbage=4, deflate=True)
    doc.close()
    return result


def csv_report(report):
    stream = StringIO(newline='')
    writer = csv.writer(stream)
    if report.get('period_summary'):
        writer.writerow([report['title'],report['description']])
        for key,label in (('completed_tasks','完成任务'),('completion_events','完成事件'),('created_tasks','新增任务'),('focus_seconds','专注秒数')):
            writer.writerow([label,report['period_summary'][key]])
        writer.writerow(['日期','新增任务','完成事件','专注秒数'])
        for day in report['daily']:writer.writerow([day['date'],day['created'],day['completed'],day['focus_seconds']])
        writer.writerow([])
    writer.writerow([label for _, label in HEADERS])
    for row in report['rows']:
        values = [display(row[key]) for key, _ in HEADERS]
        # CSV has no explicit text-cell type: neutralize spreadsheet formulas in user input.
        writer.writerow(["'" + value if value.lstrip().startswith(('=', '+', '-', '@')) else value for value in values])
    return stream.getvalue().encode('utf-8-sig')


@router.get('/exports/report')
def download_report(format: Literal['xlsx', 'pdf', 'docx', 'csv'] = 'xlsx',
                    filters: Filters = Depends(report_filters), user: User = Depends(current_user),
                    db: Session = Depends(database, scope='function')):
    report = select_report(db, user, filters)
    if format in ('pdf', 'docx') and sum(sum(len(display(value)) for value in row.values()) for row in report['rows']) > 2_000_000:
        raise HTTPException(422, '文档内容过多，请缩小范围、取消导出备注或使用 Excel')
    generators = {'xlsx': (excel_report, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                  'docx': (word_report, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
                  'pdf': (pdf_report, 'application/pdf'), 'csv': (csv_report, 'text/csv')}
    generator, media = generators[format]
    filename = f"todo-report-{utcnow().astimezone(ZoneInfo(user.timezone)):%Y-%m-%d}.{format}"
    return Response(generator(report), media_type=media,
                    headers={'Content-Disposition': f'attachment; filename="{filename}"', 'Cache-Control': 'no-store'})
