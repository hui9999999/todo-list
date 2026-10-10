import hashlib
import json
from io import BytesIO
from datetime import date, datetime, timezone
from zoneinfo import ZoneInfo
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session
from .auth import current_user
from .db import database
from .models import User, Task, ImportBatch, ImportMapping, TaskEvent, TaskSnapshot, DailyPlanItem, uid, utcnow
from .schemas import ImportRequest, TaskFields
from .tasks import lock_user, task_json, record
from .advanced_models import Project, TaskDependency
from .ai_schemas import ProjectInput
from .common import serialize
from .reports import router as reports_router

router = APIRouter()
router.include_router(reports_router)
KNOWN_V1 = {'id', 'text', 'notes', 'parentId', 'status', 'done', 'prio', 'due', 'completedAt', 'created',
            'failReason', 'hidden', 'period', 'timer', 'expanded', 'dueOverride', 'failAt'}


@router.get('/exports.xlsx')
def export_excel(user: User = Depends(current_user), db: Session = Depends(database, scope='function')):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    book = Workbook()
    sheet = book.active
    sheet.title = '任务清单'
    sheet.append(['任务 ID', '标题', '范围', '状态', '优先级', '截止日期', '截止时间', '估时（分钟）',
                  '项目', '父任务 ID', '备注', '完成时间', '垃圾站'])
    projects = {p.id: p.title for p in db.scalars(select(Project).where(Project.user_id == user.id))}
    zone = ZoneInfo(user.timezone)
    def stamp(value):
        return value.astimezone(zone).isoformat() if value else ''
    for task in db.scalars(select(Task).where(Task.user_id == user.id).order_by(Task.created_at, Task.id)):
        values = [task.id, task.title, task.scope, task.status, task.priority,
                      task.due_date, stamp(task.due_at), task.estimated_minutes,
                      projects.get(task.project_id, ''), task.parent_id or '', task.notes, stamp(task.completed_at),
                      '是' if task.deleted_at else '否']
        sheet.append([ILLEGAL_CHARACTERS_RE.sub('\ufffd', v) if isinstance(v,str) else v for v in values])
        # Explicit string cells prevent spreadsheet formula execution without altering user text.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
        sheet.cell(sheet.max_row, 6).number_format = 'yyyy-mm-dd'
    sheet.freeze_panes = 'B2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='374A43')
    for column, width in {'A':38,'B':40,'C':16,'D':16,'E':12,'F':16,'G':30,'H':18,'I':24,'J':38,'K':55,'L':30,'M':12}.items():
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    book.save(output)
    return Response(output.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': 'attachment; filename="todo-v2.xlsx"'})


def timestamp(value, zone: str):
    if value is None or value == '':
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        result = datetime.fromtimestamp(value / 1000, timezone.utc)
    elif isinstance(value, str):
        result = datetime.fromisoformat(value.replace('Z', '+00:00'))
        if result.tzinfo is None:
            tz = ZoneInfo(zone)
            candidates = [result.replace(tzinfo=tz, fold=fold) for fold in (0, 1)]
            valid = [candidate for candidate in candidates if candidate.astimezone(timezone.utc).astimezone(tz).replace(tzinfo=None) == result]
            if not valid or len({candidate.utcoffset() for candidate in valid}) > 1:
                raise ValueError('当地时间不存在或有歧义，请提供 UTC 偏移量')
            result = valid[0]
    else:
        raise ValueError('时间格式无效')
    return result.astimezone(timezone.utc).isoformat()


def backup_snapshot(value: dict, mapping: dict, user: User) -> dict:
    if not isinstance(value, dict) or value.get('id') not in mapping:
        raise ValueError('历史快照的任务引用无效')
    fields = TaskFields.model_validate({k: value[k] for k in TaskFields.model_fields if k in value})
    if fields.parent_id and fields.parent_id not in mapping:
        raise ValueError('历史快照的父任务引用缺失')
    snapshot = fields.model_dump(mode='json')
    snapshot['id'] = value['id']
    version = value.get('version')
    if not isinstance(version, int) or isinstance(version, bool) or version < 1:
        raise ValueError('历史快照的版本无效')
    snapshot['version'] = version
    for key in ('completed_at', 'created_at', 'updated_at', 'legacy_created_at', 'deleted_at'):
        snapshot[key] = timestamp(value.get(key), user.timezone)
    for key in ('legacy_period', 'deletion_group'):
        snapshot[key] = value.get(key)
    for key, default in (('timer_minutes', 25), ('timer_remaining', 1500)):
        number = value.get(key, default)
        if not isinstance(number, int) or isinstance(number, bool) or not 0 <= number <= 86400:
            raise ValueError('历史快照的计时数据无效')
        snapshot[key] = number
    return snapshot


def normalize_import(content: str, scope: str, user: User):
    try:
        payload = json.loads(content)
    except (ValueError, RecursionError):
        raise HTTPException(422, '请选择有效的 JSON 文件') from None
    is_v2 = isinstance(payload, dict) and payload.get('schema_version') == 2
    rows = payload if isinstance(payload, list) else payload.get('tasks') if isinstance(payload, dict) else None
    if not isinstance(rows, list) or len(rows) > 10000:
        raise HTTPException(422, '文件应包含任务数组，最多 10000 条')
    if not rows:
        raise HTTPException(422, '文件中没有任务')
    issues, warnings, normalized, ids = [], [], [], set()
    evidence = {'events': [], 'daily_plans': []}
    if isinstance(payload, dict) and not is_v2:
        extras = sorted(set(payload) - {'tasks'})
        if extras:
            warnings.append('此阶段仅迁移任务数组；其他顶层字段不导入：' + ', '.join(extras)[:300])
    for index, row in enumerate(rows):
        try:
            if not isinstance(row, dict):
                raise ValueError('任务必须是对象')
            old = row.get('id')
            if not isinstance(old, (str, int)) or isinstance(old, bool) or not str(old) or len(str(old)) > 100:
                raise ValueError('任务 ID 缺失或无效')
            old = str(old)
            if old in ids:
                raise ValueError('任务 ID 重复')
            ids.add(old)
            if is_v2:
                values = {k: row[k] for k in TaskFields.model_fields if k in row}
                parent_id = values.pop('parent_id', None)
                values['scope'] = row.get('scope', scope)
                completed = timestamp(row.get('completed_at'), user.timezone)
                created = timestamp(row.get('legacy_created_at') or row.get('created_at'), user.timezone)
                period = row.get('legacy_period')
                minutes = row.get('timer_minutes', 25)
                remaining = row.get('timer_remaining', minutes * 60)
                deleted_at = timestamp(row.get('deleted_at'), user.timezone)
                version = row.get('version', 1)
                deletion_group = row.get('deletion_group')
                if not isinstance(version, int) or isinstance(version, bool) or version < 1:
                    raise ValueError('任务版本无效')
                if deletion_group is not None and (not isinstance(deletion_group, str) or len(deletion_group) > 36):
                    raise ValueError('垃圾站分组无效')
            else:
                status = row.get('status') or ('done' if row.get('done') else 'todo')
                if status == 'fail':
                    status = 'failed'
                values = {'title': row.get('text'), 'notes': row.get('notes', ''), 'status': status,
                          'priority': row.get('prio', 'mid'), 'scope': scope,
                          'hidden': row.get('hidden', False), 'fail_reason': row.get('failReason', '')}
                parent_id = row.get('parentId')
                due = row.get('due')
                if due:
                    if not isinstance(due, str):
                        raise ValueError('截止日期无效')
                    if len(due) == 10:
                        values['due_date'] = date.fromisoformat(due).isoformat()
                    else:
                        values['due_at'] = timestamp(due, user.timezone)
                        values['due_timezone'] = user.timezone
                completed = timestamp(row.get('completedAt'), user.timezone)
                created = timestamp(row.get('created'), user.timezone)
                period = row.get('period')
                timer = row.get('timer') or {}
                if not isinstance(timer, dict):
                    raise ValueError('计时数据无效')
                minutes = timer.get('min', 25)
                remaining = timer.get('remaining', minutes * 60)
                deleted_at, version, deletion_group = None, 1, None
                unknown = sorted(set(row) - KNOWN_V1)
                if unknown:
                    warnings.append(f'第 {index + 1} 条：以下字段不导入：' + ', '.join(unknown)[:300])
                if row.get('failAt') is not None:
                    warnings.append(f'第 {index + 1} 条：旧 failAt 暂不迁移，失败状态与原因会保留')
                if timer.get('running'):
                    warnings.append(f'第 {index + 1} 条：计时器将暂停，保留剩余秒数')
            if parent_id is not None and (not isinstance(parent_id, (str, int)) or isinstance(parent_id, bool)):
                raise ValueError('父任务 ID 无效')
            if period not in (None, 'day', 'week', 'month'):
                raise ValueError('旧分类 period 无效')
            if not isinstance(minutes, int) or isinstance(minutes, bool) or not 1 <= minutes <= 1440:
                raise ValueError('计时分钟数无效')
            if not isinstance(remaining, (int, float)) or isinstance(remaining, bool) or not 0 <= remaining <= 86400:
                raise ValueError('剩余秒数无效')
            fields = TaskFields.model_validate(values)
            if completed and fields.status != 'done':
                raise ValueError('非完成任务含完成时间，请修正文件后重试')
            if fields.status == 'done' and not completed:
                warnings.append(f'第 {index + 1} 条：完成日期未知，不会补造历史')
            normalized.append({'legacy_id': old, 'legacy_parent_id': str(parent_id) if parent_id is not None else None,
                'fields': fields.model_dump(mode='json'), 'completed_at': completed, 'legacy_created_at': created,
                'legacy_period': period, 'timer_minutes': minutes, 'timer_remaining': int(remaining), 'deleted_at': deleted_at,
                'version': version, 'deletion_group': deletion_group})
        except (ValueError, TypeError, OverflowError, OSError, ValidationError) as error:
            # Never echo input values (they can contain credentials).
            # Do not return parser errors with raw input embedded in the message.
            controlled = {'任务必须是对象', '任务 ID 缺失或无效', '任务 ID 重复', '截止日期无效', '计时数据无效',
                          '父任务 ID 无效', '旧分类 period 无效', '计时分钟数无效', '剩余秒数无效',
                          '非完成任务含完成时间，请修正文件后重试', '任务版本无效', '垃圾站分组无效'}
            reason = str(error) if str(error) in controlled else '字段格式无效，请检查标题、状态、优先级、日期及备注'
            issues.append({'row': index + 1, 'message': reason})
    mapping = {r['legacy_id']: r for r in normalized}
    for index, item in enumerate(normalized):
        parent = item['legacy_parent_id']
        if parent and parent not in mapping:
            issues.append({'row': index + 1, 'message': '父任务缺失；请补齐父任务或清空 parentId'})
        seen, current = set(), item
        while current:
            if current['legacy_id'] in seen:
                issues.append({'row': index + 1, 'message': '父子关系存在循环，请修正文件后重试'})
                break
            seen.add(current['legacy_id'])
            current = mapping.get(current['legacy_parent_id'])
    if is_v2:
        try:
            projects = payload.get('projects', [])
            if not isinstance(projects, list) or len(projects) > 1000:
                raise ValueError()
            project_ids = set()
            evidence['projects'] = []
            for project in projects:
                identifier = project['id']
                if not isinstance(identifier, str) or len(identifier) > 36 or identifier in project_ids:
                    raise ValueError()
                project_ids.add(identifier)
                values = ProjectInput.model_validate({k: project[k] for k in ProjectInput.model_fields if k in project})
                evidence['projects'].append({'id': identifier, **values.model_dump()})
            for row in normalized:
                if row['fields'].get('project_id') and row['fields']['project_id'] not in project_ids:
                    raise ValueError()
            dependencies = payload.get('dependencies', [])
            if not isinstance(dependencies, list) or len(dependencies) > 20000:
                raise ValueError()
            evidence['dependencies'], graph = [], {}
            for dep in dependencies:
                task_id, depends_on = dep['task_id'], dep['depends_on_id']
                if task_id not in mapping or depends_on not in mapping or task_id == depends_on:
                    raise ValueError()
                graph.setdefault(task_id, set()).add(depends_on)
                evidence['dependencies'].append({'task_id': task_id, 'depends_on_id': depends_on})
            for root in graph:
                stack, seen = list(graph[root]), set()
                while stack:
                    node = stack.pop()
                    if node == root:
                        raise ValueError()
                    if node not in seen:
                        seen.add(node)
                        stack.extend(graph.get(node, []))
        except (ValueError, KeyError, TypeError, ValidationError):
            issues.append({'row': 0, 'message': '项目或依赖引用无效、重复或存在循环'})
        events, plans = payload.get('events', []), payload.get('daily_plans', [])
        if not isinstance(events, list) or not isinstance(plans, list) or len(events) > 50000 or len(plans) > 20000:
            issues.append({'row': 0, 'message': '备份历史或安排数量／格式无效'})
        else:
            for index, event in enumerate(events):
                try:
                    if not isinstance(event, dict) or event.get('task_id') not in mapping:
                        raise ValueError()
                    kind = event.get('kind')
                    if kind not in ('created', 'updated', 'status_changed', 'trashed', 'restored', 'imported', 'imported_completion', 'dependency_changed'):
                        raise ValueError()
                    occurred = timestamp(event.get('occurred_at'), user.timezone)
                    if not occurred:
                        raise ValueError()
                    after = backup_snapshot(event.get('after'), mapping, user)
                    before = backup_snapshot(event['before'], mapping, user) if event.get('before') else None
                    if after['id'] != event['task_id'] or (before and before['id'] != event['task_id']):
                        raise ValueError()
                    evidence['events'].append({'task_id': event['task_id'], 'kind': kind, 'occurred_at': occurred, 'before': before, 'after': after})
                except (ValueError, TypeError, OverflowError, OSError, ValidationError):
                    issues.append({'row': 0, 'message': f'历史第 {index + 1} 条的字段或任务引用无效'})
            seen_plans = set()
            for index, plan in enumerate(plans):
                try:
                    if not isinstance(plan, dict) or plan.get('task_id') not in mapping:
                        raise ValueError()
                    plan_date = date.fromisoformat(plan['plan_date']).isoformat()
                    key = (plan['task_id'], plan_date)
                    if key not in seen_plans:
                        seen_plans.add(key)
                        evidence['daily_plans'].append({'task_id': plan['task_id'], 'plan_date': plan_date})
                except (ValueError, TypeError, KeyError):
                    issues.append({'row': 0, 'message': f'安排第 {index + 1} 条的日期或任务引用无效'})
    return normalized, issues, warnings, evidence


def batch_json(batch: ImportBatch):
    return {'id': batch.id, 'count': len(batch.normalized), 'issues': batch.issues, 'warnings': batch.warnings,
            'confirmed': batch.confirmed_at is not None, 'items': batch.normalized,
            'status_counts': {s: sum(r['fields']['status'] == s for r in batch.normalized) for s in ('todo', 'doing', 'done', 'failed', 'cancelled')},
            'children': sum(r['legacy_parent_id'] is not None for r in batch.normalized),
            'history_count': len(batch.evidence.get('events', [])), 'plan_count': len(batch.evidence.get('daily_plans', []))}


@router.post('/imports/preview')
def preview(body: ImportRequest, user: User = Depends(current_user), db: Session = Depends(database, scope='function')):
    lock_user(db, user)
    file_digest = hashlib.sha256(body.content.encode()).hexdigest()
    existing = db.scalar(select(ImportBatch).where(ImportBatch.user_id == user.id, ImportBatch.digest == file_digest,
                                                 ImportBatch.scope == body.scope))
    if existing:
        return batch_json(existing)
    # A confirmed file cannot create copies by merely changing its default scope.
    confirmed = db.scalar(select(ImportBatch).where(ImportBatch.user_id == user.id, ImportBatch.digest == file_digest,
                                                   ImportBatch.confirmed_at.is_not(None)))
    if confirmed:
        return batch_json(confirmed)
    normalized, issues, warnings, evidence = normalize_import(body.content, body.scope, user)
    batch = ImportBatch(user_id=user.id, digest=file_digest, scope=body.scope, normalized=normalized, issues=issues, warnings=warnings, evidence=evidence)
    db.add(batch)
    db.flush()
    return batch_json(batch)


@router.post('/imports/{batch_id}/confirm')
def confirm(batch_id: str, user: User = Depends(current_user), db: Session = Depends(database, scope='function')):
    lock_user(db, user)
    batch = db.scalar(select(ImportBatch).where(ImportBatch.id == batch_id, ImportBatch.user_id == user.id).with_for_update())
    if not batch:
        raise HTTPException(404, '导入批次不存在')
    if batch.issues:
        raise HTTPException(422, '文件有异常，请修正后重新预览')
    if batch.confirmed_at:
        return {'count': len(batch.normalized), 'reused': True}
    confirmed = db.scalar(select(ImportBatch).where(ImportBatch.user_id == user.id, ImportBatch.digest == batch.digest,
                                                   ImportBatch.confirmed_at.is_not(None)))
    if confirmed:
        return {'count': len(confirmed.normalized), 'reused': True}
    id_map = {r['legacy_id']: uid() for r in batch.normalized}
    project_map = {p['id']: uid() for p in batch.evidence.get('projects', [])}
    for source in batch.evidence.get('projects', []):
        db.add(Project(id=project_map[source['id']], user_id=user.id, title=source['title'], scope=source['scope']))
    db.flush()
    created = []
    groups = {}
    # Insert all nodes first; then link parents, supporting child-before-parent arrays.
    for row in batch.normalized:
        fields = TaskFields.model_validate(row['fields']).model_dump()
        fields['parent_id'] = None
        fields['project_id'] = project_map.get(fields.get('project_id'))
        task = Task(id=id_map[row['legacy_id']], user_id=user.id, **fields,
                    completed_at=datetime.fromisoformat(row['completed_at']) if row['completed_at'] else None,
                    legacy_created_at=datetime.fromisoformat(row['legacy_created_at']) if row['legacy_created_at'] else None,
                    legacy_period=row['legacy_period'], timer_minutes=row['timer_minutes'], timer_remaining=row['timer_remaining'],
                    deleted_at=datetime.fromisoformat(row['deleted_at']) if row['deleted_at'] else None,
                    deletion_group=groups.setdefault(row['deletion_group'] or row['legacy_id'], uid()) if row['deleted_at'] else None,
                    version=row['version'])
        db.add(task)
        created.append((row, task))
    db.flush()
    history_task_ids = set()
    def remap(snapshot):
        if snapshot is None:
            return None
        mapped = dict(snapshot)
        mapped['id'] = id_map[mapped['id']]
        mapped['parent_id'] = id_map.get(mapped['parent_id'])
        mapped['project_id'] = project_map.get(mapped.get('project_id'))
        mapped['deletion_group'] = groups.get(mapped.get('deletion_group'))
        return mapped
    for source in batch.evidence.get('events', []):
        history_task_ids.add(source['task_id'])
        before, after = remap(source['before']), remap(source['after'])
        event = TaskEvent(user_id=user.id, task_id=id_map[source['task_id']], kind=source['kind'], actor='import',
                          occurred_at=datetime.fromisoformat(source['occurred_at']),
                          before_version=before['version'] if before else None, after_version=after['version'])
        db.add(event)
        db.flush()
        db.add(TaskSnapshot(user_id=user.id, event_id=event.id, before=before, after=after))
    for source in batch.evidence.get('daily_plans', []):
        db.add(DailyPlanItem(user_id=user.id, task_id=id_map[source['task_id']], plan_date=date.fromisoformat(source['plan_date'])))
    seen_dependencies = set()
    for source in batch.evidence.get('dependencies', []):
        key = (source['task_id'], source['depends_on_id'])
        if key not in seen_dependencies:
            seen_dependencies.add(key)
            db.add(TaskDependency(user_id=user.id, task_id=id_map[key[0]], depends_on_id=id_map[key[1]]))
    for row, task in created:
        task.parent_id = id_map.get(row['legacy_parent_id'])
    db.flush()
    for row, task in created:
        record(db, task, 'imported')
        if task.completed_at and row['legacy_id'] not in history_task_ids:
            record(db, task, 'imported_completion', occurred_at=task.completed_at)
        db.add(ImportMapping(user_id=user.id, batch_id=batch.id, legacy_id=row['legacy_id'], task_id=task.id))
    batch.confirmed_at = utcnow()
    # Minimize duplicated task bodies once import is committed.
    batch.normalized = [{'legacy_id': r['legacy_id'], 'legacy_parent_id': r['legacy_parent_id'],
                         'fields': {'status': r['fields']['status']}} for r in batch.normalized]
    batch.evidence = {}
    return {'count': len(created), 'reused': False}


@router.get('/exports')
def export(compat: bool = False, user: User = Depends(current_user), db: Session = Depends(database, scope='function')):
    tasks = list(db.scalars(select(Task).where(Task.user_id == user.id).order_by(Task.created_at, Task.id)))
    if compat:
        zone = ZoneInfo(user.timezone)
        def ms(dt):
            return int(dt.timestamp() * 1000) if dt else None
        return [{'id': t.id, 'text': t.title, 'notes': t.notes, 'parentId': t.parent_id,
                 'status': 'fail' if t.status == 'failed' else 'todo' if t.status in ('doing', 'cancelled') else t.status,
                 'done': t.status == 'done', 'prio': t.priority, 'failReason': t.fail_reason, 'hidden': t.hidden,
                 'due': t.due_date.isoformat() if t.due_date else t.due_at.astimezone(zone).strftime('%Y-%m-%dT%H:%M') if t.due_at else '',
                 'period': t.legacy_period or 'day', 'created': ms(t.legacy_created_at or t.created_at),
                 'completedAt': ms(t.completed_at), 'timer': {'min': t.timer_minutes, 'remaining': t.timer_remaining, 'running': False}}
                for t in tasks if not t.deleted_at]
    rows = db.execute(select(TaskEvent, TaskSnapshot).join(TaskSnapshot, TaskSnapshot.event_id == TaskEvent.id)
                      .where(TaskEvent.user_id == user.id).order_by(TaskEvent.recorded_at))
    return {'schema_version': 2, 'exported_at': utcnow().isoformat(), 'timezone': user.timezone,
            'projects': [serialize(p) for p in db.scalars(select(Project).where(Project.user_id == user.id))],
            'dependencies': [serialize(d) for d in db.scalars(select(TaskDependency).where(TaskDependency.user_id == user.id))],
            'tasks': [task_json(t) for t in tasks],
            'events': [{'id': e.id, 'task_id': e.task_id, 'kind': e.kind, 'actor': e.actor, 'occurred_at': e.occurred_at.isoformat(),
                        'recorded_at': e.recorded_at.isoformat(), 'before': s.before, 'after': s.after} for e, s in rows],
            'daily_plans': [{'task_id': p.task_id, 'plan_date': p.plan_date.isoformat()} for p in db.scalars(
                select(DailyPlanItem).where(DailyPlanItem.user_id == user.id))]}
