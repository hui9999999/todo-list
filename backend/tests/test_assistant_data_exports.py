import json
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
import pytest
from cryptography.fernet import Fernet
from pydantic import ValidationError
from sqlalchemy import select
from app.common import tenant_session, today
from app.models import Task, User, utcnow
from app.advanced_models import ProviderConnection, UsageEvent, FocusSession, Message, Memory, RunEvent
from app.config import settings
from app.assistant_schemas import DataQuery, ExportRequest
from app.assistant_data import query_data, query_evidence, account_archive
from app.ai_schemas import AgentOutput
from test_p1_p4 import connection, run_message, output_stream, task
from test_ai_task_operations import json_stream

P='/api/v1'


@pytest.fixture(autouse=True)
def encryption(monkeypatch):
    monkeypatch.setattr(settings(),'credential_master_key',Fernet.generate_key().decode())


def execute_export(client,monkeypatch,request):
    import app.agent as agent
    monkeypatch.setattr(agent,'stream_chat',lambda c,p:output_stream({'answer':'文件已准备好，请点击下载。','exports':[request]}))
    chat,run,_=run_message(client,'请导出我的资料')
    agent.execute_agent(client.user_id,run['id'])
    info=client.get(P+'/agent-runs/'+run['id']).json()
    assert info['status']=='completed',info
    data=client.get(P+'/conversations/'+chat['id']+'/messages').json()
    assert len(data['downloads'])==1
    card=data['downloads'][0]
    assert card['message_id']==data['messages'][-1]['id']
    result=client.get(P+'/assistant/exports/'+card['id'])
    assert result.status_code==200,result.text[:300] if result.status_code!=200 else ''
    return chat,card,result


def test_all_categories_include_hidden_trash_and_account_without_credentials(accounts):
    client,other=accounts(),accounts()
    connection(client)
    hidden=task(client,'隐藏的个人事项',hidden=True)
    discarded=task(client,'垃圾站个人事项')
    client.post(P+f"/tasks/{discarded['id']}/trash",json={'version':discarded['version']})
    secret=task(other,'其他用户事项禁止读取')
    with tenant_session(client.user_id) as db:
        user=db.get(User,client.user_id)
        profile=query_data(db,user,DataQuery(category='account'))['items'][0]
        assert profile['email']==client.email and 'password_hash' not in profile
        overview=query_data(db,user,DataQuery(category='all'))
        assert next(c for c in overview['categories'] if c['category']=='trash')['total']==1
        assert query_data(db,user,DataQuery(category='tasks',id=hidden['id']))['items'][0]['hidden']
        assert query_data(db,user,DataQuery(category='tasks',id=secret['id']))['total']==0
        archive=account_archive(db,user)
        encoded=json.dumps(archive,ensure_ascii=False,default=str)
        assert discarded['id'] in encoded and hidden['id'] in encoded
        assert '其他用户事项' not in encoded
        assert not any(key in encoded for key in ('password_hash','ciphertext','csrf_token','test-key-never-export'))
        assert all(set(s)=={'expires_at'} for s in archive['sessions'])


def test_query_pages_past_initial_100_and_escapes_like_patterns(accounts):
    client=accounts()
    with tenant_session(client.user_id) as db:
        user=db.get(User,client.user_id)
        db.add_all(Task(user_id=user.id,title=f'长清单-{i:03}',notes='',scope='study') for i in range(125))
        db.add(Task(user_id=user.id,title='百分号%与下划线_',notes=''))
        db.flush()
        result=query_data(db,user,DataQuery(category='tasks',q='长清单',offset=100,limit=20))
        assert result['total']==125 and len(result['items'])==20 and result['next_offset']==120
        assert query_data(db,user,DataQuery(category='tasks',q='%'))['total']==1
        narrowed=query_data(db,user,DataQuery(category='tasks',q='长清单',fields=['title'],limit=1))
        assert set(narrowed['items'][0])=={'id','title'}
        from fastapi import HTTPException
        with pytest.raises(HTTPException):query_data(db,user,DataQuery(category='connections',fields=['ciphertext']))
        tiny=query_evidence(db,user,DataQuery(category='tasks',q='长清单'),budget=1)
        assert tiny['total']==125 and tiny['next_offset']==1 and tiny['items'][0]['content_omitted']


def test_query_due_dates_use_account_timezone_and_include_end_date(accounts):
    client=accounts()
    task(client,'当天日期',due_date='2026-10-09')
    task(client,'当天最后秒',due_at='2026-10-09T23:59:59+08:00',due_timezone='Asia/Shanghai')
    task(client,'下一天',due_at='2026-10-10T00:00:00+08:00',due_timezone='Asia/Shanghai')
    with tenant_session(client.user_id) as db:
        result=query_data(db,db.get(User,client.user_id),DataQuery(category='tasks',date_field='due',start=date(2026,10,9),end=date(2026,10,9)))
        assert result['total']==2


@pytest.mark.parametrize('tools',[True,False])
@pytest.mark.parametrize('query_format',['structured','content_json','prose_json','answer_json'])
def test_assistant_queries_real_trash_then_answers_in_both_transports(accounts,monkeypatch,tools,query_format):
    import app.agent as agent
    client,other=accounts(),accounts();conn=connection(client)
    trashed=task(client,'被删的学习计划',hidden=True)
    client.post(P+f"/tasks/{trashed['id']}/trash",json={'version':trashed['version']})
    other_task=task(other,'其他用户的秘密')
    other.post(P+f"/tasks/{other_task['id']}/trash",json={'version':other_task['version']})
    with tenant_session(client.user_id) as db:db.get(ProviderConnection,conn['id']).capabilities={'test-model':{'chat':True,'tools':tools}}
    calls=[]
    def respond(c,payload):
        calls.append(payload)
        if len(calls)==1:
            output={'answer':'我查一下你的垃圾站。','queries':[{'category':'trash'}]}
            if query_format=='answer_json':
                output={'answer':'我试试查一下垃圾站。\n\n'+json.dumps({'queries':[{'category':'trash','limit':50}]},ensure_ascii=False)}
            elif query_format!='structured':
                content=json.dumps(output,ensure_ascii=False) if query_format=='content_json' else '我试试查一下垃圾站。\n\n'+json.dumps({'queries':[{'category':'trash','limit':50}]},ensure_ascii=False)
                for start in range(0,len(content),7):
                    yield {'choices':[{'delta':{'content':content[start:start+7]}}]}
                yield {'choices':[{'delta':{},'finish_reason':'stop'}],'usage':{'prompt_tokens':30,'completion_tokens':20}}
                return
        else:
            evidence=payload['messages'][-1]['content']
            assert '被删的学习计划' in evidence and trashed['id'] in evidence
            assert '其他用户的秘密' not in evidence and 'ciphertext' not in evidence
            output={'answer':'垃圾站中有被删的学习计划，当前状态为待办。'}
        yield from (output_stream if tools else json_stream)(output)
    monkeypatch.setattr(agent,'stream_chat',respond)
    chat,run,_=run_message(client,'我的垃圾站中有什么？')
    agent.execute_agent(client.user_id,run['id'])
    info=client.get(P+'/agent-runs/'+run['id']).json()
    assert info['status']=='completed' and info['answer'].startswith('垃圾站中有') and not info['proposals']
    assert len(calls)==2
    messages=client.get(P+'/conversations/'+chat['id']+'/messages').json()['messages']
    assert messages[-1]['content']==info['answer'] and len(messages)==3
    with tenant_session(client.user_id) as db:
        usages=list(db.scalars(select(UsageEvent).where(UsageEvent.user_id==client.user_id,UsageEvent.purpose=='chat')))
        assert len(usages)==2 and all(u.state=='settled' for u in usages)
        events=list(db.scalars(select(RunEvent).where(RunEvent.user_id==client.user_id,RunEvent.run_id==run['id'],RunEvent.kind=='answer_delta')))
        assert all('"queries"' not in event.payload['text'] and '"category"' not in event.payload['text'] for event in events)


def test_invalid_or_unbounded_queries_never_write_tasks(accounts,monkeypatch):
    import app.agent as agent
    with pytest.raises(ValidationError):DataQuery.model_validate({'category':'tasks','user_id':'another'})
    with pytest.raises(ValidationError):DataQuery.model_validate({'category':'passwords'})
    with pytest.raises(ValidationError):AgentOutput.model_validate({'answer':'操作','queries':[{'category':'tasks'}],'actions':[{'kind':'create','fields':{'title':'不得执行'}}]})
    client=accounts();connection(client)
    monkeypatch.setattr(agent,'stream_chat',lambda c,p:output_stream({'answer':'继续查找','queries':[{'category':'tasks'}]}))
    _,run,_=run_message(client,'查询资料')
    agent.execute_agent(client.user_id,run['id'])
    info=client.get(P+'/agent-runs/'+run['id']).json()
    assert info['status']=='failed' and '查询范围较大' in info['error']
    assert client.get(P+'/tasks').json()==[]


@pytest.mark.parametrize('tools',[True,False])
def test_trash_followup_has_current_records_despite_old_permission_claim(accounts,monkeypatch,tools):
    import app.agent as agent
    client,other=accounts(),accounts();conn=connection(client)
    discarded=task(client,'垃圾站中的真实学习事项')
    client.post(P+f"/tasks/{discarded['id']}/trash",json={'version':discarded['version']})
    foreign=task(other,'别人的垃圾站不得读取')
    other.post(P+f"/tasks/{foreign['id']}/trash",json={'version':foreign['version']})
    chat=client.post(P+'/conversations').json()
    with tenant_session(client.user_id) as db:
        db.get(ProviderConnection,conn['id']).capabilities={'test-model':{'chat':True,'tools':tools}}
        db.add(Message(user_id=client.user_id,conversation_id=chat['id'],role='assistant',
            content='我没有垃圾站读取权限。\n{"queries":[{"category":"trash","limit":50}]}',created_at=utcnow()-timedelta(minutes=1)))
    run=client.post(P+'/conversations/'+chat['id']+'/messages',json={'content':'现在能看了吗','idempotency_key':'trash-followup-current-data'}).json()
    def respond(c,payload):
        system=payload['messages'][0]['content']
        assert '垃圾站中的真实学习事项' in system and discarded['id'] in system
        assert '别人的垃圾站不得读取' not in system
        assert '系统已为当前垃圾站话题执行只读查询' in system
        yield from (output_stream if tools else json_stream)({'answer':'垃圾站里有垃圾站中的真实学习事项。'})
    monkeypatch.setattr(agent,'stream_chat',respond)
    agent.execute_agent(client.user_id,run['id'])
    info=client.get(P+'/agent-runs/'+run['id']).json()
    assert info['status']=='completed' and not info['proposals'] and '真实学习事项' in info['answer']
    assert client.get(P+'/tasks?trash=true').json()[0]['version']==discarded['version']+1
    with tenant_session(client.user_id) as db:
        events=list(db.scalars(select(RunEvent).where(RunEvent.user_id==client.user_id,RunEvent.run_id==run['id'],RunEvent.kind=='tool_result')))
        assert any(event.payload=={'kind':'read_data','categories':['trash']} for event in events)


@pytest.mark.parametrize('kind',['weekly','monthly'])
@pytest.mark.parametrize('format',['xlsx','docx','pdf','csv'])
def test_calendar_reports_export_real_events_and_focus(accounts,monkeypatch,kind,format):
    client=accounts();connection(client)
    completed=task(client,'本期已完成学习',scope='study',status='done')
    with tenant_session(client.user_id) as db:
        user=db.get(User,client.user_id);day=today(user)
        db.add(FocusSession(user_id=user.id,task_id=completed['id'],state='finished',duration_seconds=600,elapsed_seconds=600,started_at=utcnow(),anchor_at=utcnow()))
    task(client,'本期待结转',scope='work',due_date=str(day))
    _,card,response=execute_export(client,monkeypatch,{'kind':kind,'format':format,'date':str(day)})
    assert card['format']==format and '完成 1 项' in card['description'] and '专注 10.0 分钟' in card['description']
    title='周报' if kind=='weekly' else '月报'
    if format=='xlsx':
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(response.content))
        assert book['周期统计']['A1'].value.endswith(title)
        assert book['周期统计']['B2'].value==1 and book['周期统计']['B5'].value==600
        assert sum(row[3] or 0 for row in list(book['每日统计'].values)[1:])==600
        assert len(list(book['任务清单'].values))==3
    elif format=='docx':
        from docx import Document
        text='\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs)
        assert title in text and '专注 10.0 分钟' in text and '本期已完成学习' in text
    elif format=='pdf':
        import pymupdf
        doc=pymupdf.open(stream=response.content,filetype='pdf')
        text=''.join(page.get_text() for page in doc)
        assert title in text and '本期已完成学习' in text and '10.0' in text
    else:
        text=response.content.decode('utf-8-sig')
        assert title in text and '专注秒数,600' in text and '本期已完成学习' in text and '本期待结转' in text


def test_account_export_download_is_private_and_survives_reload_until_chat_deleted(accounts,monkeypatch):
    client,other=accounts(),accounts();connection(client)
    original=task(client,'个人资料内容',hidden=True)
    chat,card,response=execute_export(client,monkeypatch,{'kind':'account','format':'json'})
    archive=response.json()
    assert archive['account']['email']==client.email and original['id'] in json.dumps(archive)
    assert all('csrf_token' not in row for row in archive['sessions'])
    assert all('ciphertext' not in row for row in archive['connections'])
    assert response.headers['cache-control']=='no-store'
    assert other.get(P+'/assistant/exports/'+card['id']).status_code==404
    assert client.get(P+'/conversations/'+chat['id']+'/messages').json()['downloads'][0]==card
    assert client.delete(P+'/conversations/'+chat['id']).status_code==200
    assert client.get(P+'/assistant/exports/'+card['id']).status_code==404


def test_task_export_filter_and_formula_safety(accounts,monkeypatch):
    client=accounts();connection(client)
    discarded=task(client,'=2+3',scope='study')
    client.post(P+f"/tasks/{discarded['id']}/trash",json={'version':discarded['version']})
    task(client,'不应导出的工作',scope='work')
    _,_,response=execute_export(client,monkeypatch,{'kind':'tasks','format':'csv','scope':'study','include_trash':True})
    text=response.content.decode('utf-8-sig')
    assert "'=2+3" in text and '不应导出的工作' not in text


def test_forgotten_sources_and_downloads_are_hidden_before_cleanup(accounts,monkeypatch):
    client=accounts();connection(client)
    chat,card,_=execute_export(client,monkeypatch,{'kind':'account','format':'json'})
    with tenant_session(client.user_id) as db:
        source=db.scalar(select(Message).where(Message.user_id==client.user_id,Message.conversation_id==chat['id'],Message.role=='user'))
        memory=Memory(user_id=client.user_id,content='将遗忘的内容',key='遗忘',source_message_id=source.id)
        db.add(memory);db.flush();identifier=memory.id
    assert client.post(P+'/memories/'+identifier+'/forget').status_code==200
    assert client.get(P+'/assistant/exports/'+card['id']).status_code==404
    with tenant_session(client.user_id) as db:
        messages=query_data(db,db.get(User,client.user_id),DataQuery(category='messages',conversation_id=chat['id']))
        assert messages['total']==0
        conversation=query_data(db,db.get(User,client.user_id),DataQuery(category='conversations',id=chat['id']))
        assert conversation['items'][0]['title']=='[来源已遗忘的对话]'


def test_export_rejects_future_period_and_invalid_ranges(accounts,monkeypatch):
    import app.agent as agent
    client=accounts();connection(client)
    monkeypatch.setattr(agent,'stream_chat',lambda c,p:output_stream({'answer':'生成报告','exports':[{'kind':'weekly','date':'2099-01-01'}]}))
    chat,run,_=run_message(client,'导出未来周报')
    agent.execute_agent(client.user_id,run['id'])
    assert '尚未开始' in client.get(P+'/agent-runs/'+run['id']).json()['error']
    assert client.get(P+'/conversations/'+chat['id']+'/messages').json()['downloads']==[]
    with pytest.raises(ValidationError):ExportRequest(kind='weekly',scope='work')
    with pytest.raises(ValidationError):ExportRequest(kind='tasks',start=date(2026,10,10),end=date(2026,10,1))
