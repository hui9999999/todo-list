from datetime import date, datetime, timezone
from io import BytesIO, StringIO
from pathlib import Path
import csv
import re
import pytest
from app.common import tenant_session
from app.models import Task, User

PREFIX = '/api/v1'


def task(client, title, **fields):
    result = client.post(PREFIX + '/tasks', json={'title': title, **fields})
    assert result.status_code == 201, result.text
    return result.json()


def set_dates(client, identifier, **fields):
    with tenant_session(client.user_id) as db:
        row = db.get(Task, identifier)
        for key, value in fields.items():
            setattr(row, key, value)


def test_report_filters_timezone_boundaries_and_owner(accounts):
    client, other = accounts(), accounts()
    first = task(client, '日期起点', scope='study', hidden=True, status='done')
    last = task(client, '日期末尾', scope='study')
    outside = task(client, '下一天', scope='study')
    task(client, '工作对照', scope='work')
    trashed = task(client, '垃圾站学习', scope='study')
    assert client.post(PREFIX + f"/tasks/{trashed['id']}/trash", json={'version': trashed['version']}).status_code == 200
    start = datetime(2026, 10, 1, 16, tzinfo=timezone.utc)
    set_dates(client, first['id'], created_at=start, legacy_created_at=start,
              updated_at=start, completed_at=start, due_at=None, due_timezone=None, due_date=date(2026, 10, 2))
    set_dates(client, last['id'], created_at=datetime(2026, 10, 2, 15, 59, 59, tzinfo=timezone.utc),
              due_at=datetime(2026, 10, 2, 15, 59, tzinfo=timezone.utc), due_timezone='Asia/Shanghai', due_date=None)
    set_dates(client, outside['id'], created_at=datetime(2026, 10, 2, 16, tzinfo=timezone.utc), due_at=None, due_timezone=None, due_date=None)
    set_dates(client, trashed['id'], created_at=start)
    query = '?start=2026-10-02&end=2026-10-02&scope=study'
    preview = client.get(PREFIX + '/exports/preview' + query)
    assert preview.status_code == 200, preview.text
    assert preview.json()['summary']['total'] == 2
    assert preview.json()['summary']['scopes']['study'] == 2
    assert client.get(PREFIX + '/exports/preview' + query + '&include_trash=true').json()['summary']['total'] == 3
    assert client.get(PREFIX + '/exports/preview' + query + '&date_field=due').json()['summary']['total'] == 2
    assert client.get(PREFIX + '/exports/preview' + query + '&date_field=completed').json()['summary']['total'] == 1
    assert client.get(PREFIX + '/exports/preview' + query + '&date_field=updated').json()['summary']['total'] >= 1
    assert client.get(PREFIX + '/exports/preview' + query + '&status=done').json()['summary']['total'] == 1
    csv_data = other.get(PREFIX + '/exports/report?format=csv').content.decode('utf-8-sig')
    assert len(list(csv.reader(StringIO(csv_data)))) == 1
    for invalid in ('start=2026-10-03&end=2026-10-02', 'end=9999-12-31', 'date_field=invalid', 'scope=invalid', 'format=exe'):
        assert client.get(PREFIX + '/exports/report?' + invalid).status_code == 422


def test_report_dst_and_legacy_creation(accounts):
    client = accounts()
    with tenant_session(client.user_id) as db:
        db.get(User, client.user_id).timezone = 'America/New_York'
    first, last, outside = [task(client, title) for title in ('开始', '结束', '下一天')]
    set_dates(client, first['id'], legacy_created_at=datetime(2026, 3, 8, 5, tzinfo=timezone.utc))
    set_dates(client, last['id'], created_at=datetime(2026, 3, 9, 3, 59, 59, tzinfo=timezone.utc))
    set_dates(client, outside['id'], created_at=datetime(2026, 3, 9, 4, tzinfo=timezone.utc))
    result = client.get(PREFIX + '/exports/preview?start=2026-03-08&end=2026-03-08')
    assert result.status_code == 200 and result.json()['summary']['total'] == 2, result.text


def test_all_report_formats_chinese_notes_dates_and_literal_formulas(accounts):
    import pymupdf
    from docx import Document
    from openpyxl import load_workbook
    client = accounts()
    project = client.post(PREFIX + '/projects', json={'title': '学习项目', 'scope': 'study'}).json()
    item = task(client, '=1+1 中文学习任务', scope='study', status='done', project_id=project['id'],
                estimated_minutes=45, notes='备注必须完整导出\n' + '长备注换页测试内容。' * 140 + '\n' + 'ABCDEFGHIJKLMNOPQRSTUVWXYZ' * 12, due_date='2026-10-02')
    task(client, '工作计划', scope='work', priority='high', notes='工作备注')
    qa = Path('.runtime/report-qa')
    qa.mkdir(parents=True, exist_ok=True)
    for format in ('xlsx', 'docx', 'pdf', 'csv'):
        result = client.get(PREFIX + '/exports/report?format=' + format)
        assert result.status_code == 200, result.text[:300] if result.status_code != 200 else ''
        assert f'.{format}' in result.headers['content-disposition']
        assert result.headers['cache-control'] == 'no-store'
        (qa / f'sample.{format}').write_bytes(result.content)
        if format == 'xlsx':
            book = load_workbook(BytesIO(result.content))
            sheet = book['任务清单']
            assert sheet.max_row == 3 and sheet['A2'].value == item['title'] and sheet['A2'].data_type == 's'
            assert sheet['B2'].value == '学习' and sheet['F2'].data_type == 'd'
            assert sheet['K2'].value == 45 and sheet['L2'].value == item['notes']
            assert len(book['导出概览']._charts) == 2
        elif format == 'docx':
            doc = Document(BytesIO(result.content))
            text = '\n'.join(p.text for p in doc.paragraphs)
            assert item['title'] in text and item['notes'] in text and '学习项目' in text
            assert len(doc.inline_shapes) == 2
        elif format == 'pdf':
            doc = pymupdf.open(stream=result.content, filetype='pdf')
            # Exclude the page footer so a note spanning pages stays contiguous in this check.
            text = ''.join(page.get_text(clip=pymupdf.Rect(0, 0, page.rect.width, 790)) for page in doc)
            assert item['title'] in text and '备注必须完整导出' in text and '学习项目' in text
            assert re.sub(r'\s', '', item['notes']) in re.sub(r'[\s\u200b]', '', text)
            assert len(doc) >= 2
            for index, page in enumerate(doc):
                page.get_pixmap(matrix=pymupdf.Matrix(1.3, 1.3)).save(qa / f'pdf-page-{index + 1}.png')
            doc.close()
        else:
            rows = list(csv.reader(StringIO(result.content.decode('utf-8-sig'))))
            assert rows[1][0] == "'" + item['title'] and rows[1][11] == item['notes']
    without = client.get(PREFIX + '/exports/report?format=docx&include_notes=false&scope=study')
    assert without.status_code == 200
    text = '\n'.join(p.text for p in Document(BytesIO(without.content)).paragraphs)
    assert '备注必须完整导出' not in text and item['title'] in text and '工作计划' not in text
    empty = client.get(PREFIX + '/exports/report?format=pdf&scope=life')
    doc = pymupdf.open(stream=empty.content, filetype='pdf')
    assert '没有符合筛选条件的任务' in ''.join(page.get_text() for page in doc)
    doc.close()
