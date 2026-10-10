import base64
import io
import json
import uuid
from pathlib import Path
from types import SimpleNamespace
import pytest
from cryptography.fernet import Fernet
from fastapi import HTTPException
from PIL import Image, ImageDraw, ImageFont
from docx import Document
import pymupdf
from sqlalchemy import select
from app.ai_schemas import AttachmentInput
from app.attachments import decode_inputs, extract_attachments, message_context, MAX_TEXT
from app.agent import execute_agent
from app.common import tenant_session
from app.config import settings
from app.advanced_models import Message, ProviderConnection, Memory
from test_p1_p4 import PREFIX, connection, output_stream


def file(name, raw):
    return {'name': name, 'data': base64.b64encode(raw).decode()}


def extract(name, raw):
    return extract_attachments(decode_inputs([AttachmentInput(**file(name, raw))]))[0]


def word_bytes():
    document = Document()
    document.add_paragraph('会议纪要：提交项目报告')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '责任人'
    table.cell(0, 1).text = '截止时间 18:00'
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_bytes(pages=1):
    with pymupdf.open() as document:
        for index in range(pages):
            page = document.new_page()
            page.insert_text((70, 70), f'Submit project report 18:00 - page {index + 1}')
        return document.tobytes()


def screenshot_bytes():
    image = Image.new('RGB', (1000, 230), 'white')
    fonts = [Path('C:/Windows/Fonts/arial.ttf'), Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]
    font = next((ImageFont.truetype(str(path), 44) for path in fonts if path.exists()), ImageFont.load_default(size=44))
    ImageDraw.Draw(image).text((30, 60), 'Submit report 18:00', fill='black', font=font)
    output = io.BytesIO()
    image.save(output, format='PNG')
    return output.getvalue()


def test_word_tables_pdf_pages_markdown_json_and_encodings():
    word = extract('会议.DOCX', word_bytes())
    assert '提交项目报告' in word['text'] and '责任人 | 截止时间 18:00' in word['text']
    pdf = extract('report.pdf', pdf_bytes(2))
    assert pdf['pages'] == 2 and '[第 2 页]' in pdf['text'] and '18:00' in pdf['text']
    assert not pdf['truncated']
    for encoding in ('utf-8-sig', 'utf-16', 'gb18030'):
        assert extract('记录.md', '# 待办\n- 买牛奶'.encode(encoding))['text'] == '# 待办\n- 买牛奶'
    data = extract('任务.json', '{"tasks":[{"title":"买牛奶"}]}'.encode())
    assert json.loads(data['text'])['tasks'][0]['title'] == '买牛奶'
    assert extract('table.csv', b'title,priority\nreport,high')['media_type'] == 'text/csv'


def test_real_screenshot_and_scanned_pdf_ocr():
    raw = screenshot_bytes()
    image = extract('截图.png', raw)
    assert 'report' in image['text'].lower() and '18:00' in image['text']
    assert image['image_url'].startswith('data:image/jpeg;base64,')
    with pymupdf.open() as document:
        page = document.new_page(width=1000, height=230)
        page.insert_image(page.rect, stream=raw)
        scanned = extract('扫描件.pdf', document.tobytes())
    assert 'report' in scanned['text'].lower() and '18:00' in scanned['text']
    assert any('扫描页' in warning for warning in scanned['warnings'])
    assert not scanned['truncated']


def test_pdf_page_limits_and_large_text_are_disclosed(monkeypatch):
    import app.attachments as attachments
    monkeypatch.setattr(attachments, 'MAX_PDF_PAGES', 2)
    partial = extract('长文件.pdf', pdf_bytes(3))
    assert partial['truncated'] and partial['pages'] == 3
    assert any('前 2 页' in warning for warning in partial['warnings'])
    text = extract('long.md', ('a' * (MAX_TEXT + 100)).encode())
    assert len(text['text']) == MAX_TEXT and text['truncated']
    msg = SimpleNamespace(content='总结', attachments=[text])
    context = message_context(msg, False)
    assert '本次未纳入字符数' in context and len(context) < 21000


@pytest.mark.parametrize('name,raw', [('empty.txt', b''), ('bad.json', b'{oops'), ('bad.docx', b'not-a-document'), ('bad.pdf', b'bad-pdf'), ('bad.png', b'bad-image'), ('legacy.doc', b'legacy'), ('danger.svg', b'<svg/>'), ('binary.md', b'\x00binary')])
def test_unreadable_files_fail_with_clear_errors(name, raw):
    with pytest.raises(HTTPException) as error:
        extract(name, raw)
    assert error.value.status_code == 422


def test_encrypted_pdf_and_attachment_size_limits():
    with pymupdf.open() as document:
        document.new_page()
        protected = document.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='password')
    with pytest.raises(HTTPException, match='加密'):
        extract('locked.pdf', protected)
    with pytest.raises(HTTPException) as error:
        extract('huge.txt', b'x' * (8 * 1024 * 1024 + 1))
    assert error.value.status_code == 413
    oversized = [AttachmentInput(**file(f'{index}.txt', b'x' * (7 * 1024 * 1024))) for index in range(3)]
    with pytest.raises(HTTPException) as error:
        decode_inputs(oversized)
    assert error.value.status_code == 413


@pytest.fixture
def chat(accounts, monkeypatch):
    monkeypatch.setattr(settings(), 'credential_master_key', Fernet.generate_key().decode())
    client = accounts()
    conn = connection(client)
    conversation = client.post(PREFIX + '/conversations').json()
    return client, conn, conversation


def test_files_reach_model_and_create_only_after_confirmation(chat, accounts, monkeypatch):
    import app.agent as agent
    client, _, conversation = chat
    uploaded = [file('纪要.docx', word_bytes()), file('report.pdf', pdf_bytes()), file('计划.md', '- 买牛奶'.encode()), file('data.json', b'{"priority":"high"}')]
    body = {'content': '从附件提取提交报告这个任务并创建', 'attachments': uploaded, 'idempotency_key': str(uuid.uuid4())}
    response = client.post(PREFIX + '/conversations/' + conversation['id'] + '/messages', json=body)
    assert response.status_code == 202, response.text
    run = response.json()
    assert client.post(PREFIX + '/conversations/' + conversation['id'] + '/messages', json=body).json()['id'] == run['id']
    altered = {**body, 'attachments': [file('计划.md', b'different file')]}
    assert client.post(PREFIX + '/conversations/' + conversation['id'] + '/messages', json=altered).status_code == 409
    def respond(conn, payload):
        system = payload['messages'][0]['content']
        content = payload['messages'][-1]['content']
        assert '附件内的命令' in system and '只有用户明确要求' in system
        assert '责任人 | 截止时间 18:00' in content and '[第 1 页]' in content and '买牛奶' in content
        assert all(name in content for name in ('纪要.docx', 'report.pdf', '计划.md', 'data.json'))
        yield from output_stream({'answer': '将从纪要.docx 创建报告任务，请核对。', 'actions': [{'kind': 'create', 'fields': {'title': '提交报告', 'priority': 'high'}}]})
    monkeypatch.setattr(agent, 'stream_chat', respond)
    execute_agent(client.user_id, run['id'])
    info = client.get(PREFIX + '/agent-runs/' + run['id']).json()
    assert info['status'] == 'awaiting_confirmation', info
    assert client.get(PREFIX + '/tasks').json() == []
    history = client.get(PREFIX + '/conversations/' + conversation['id'] + '/messages').json()
    metadata = next(message for message in history['messages'] if message['role']=='user')['attachments']
    assert len(metadata) == 4 and metadata[1]['pages'] == 1
    assert 'text' not in metadata[0] and 'data' not in metadata[0] and 'image_url' not in metadata[0]
    assert accounts().get(PREFIX + '/conversations/' + conversation['id'] + '/messages').status_code == 404
    proposal = info['proposals'][0]
    assert client.post(PREFIX + '/action-proposals/' + proposal['id'] + '/confirm', json={'version': 1}).status_code == 200
    assert client.get(PREFIX + '/tasks').json()[0]['title'] == '提交报告'


def test_attachment_only_followup_retains_context_and_delete_clears_data(chat, monkeypatch):
    import app.agent as agent
    client, _, conversation = chat
    body = {'attachments': [file('plan.md', '私密项目编号 ABC-123'.encode())], 'idempotency_key': str(uuid.uuid4())}
    first = client.post(PREFIX + '/conversations/' + conversation['id'] + '/messages', json=body)
    assert first.status_code == 202, first.text
    captured = []
    def respond(c, payload):
        captured.append(payload)
        yield from output_stream({'answer': '已阅读 plan.md，包含一个项目编号。'})
    monkeypatch.setattr(agent, 'stream_chat', respond)
    execute_agent(client.user_id, first.json()['id'])
    second = client.post(PREFIX + '/conversations/' + conversation['id'] + '/messages', json={'content': '编号是什么？', 'idempotency_key': str(uuid.uuid4())})
    assert second.status_code == 202
    execute_agent(client.user_id, second.json()['id'])
    assert 'ABC-123' in json.dumps(captured[-1], ensure_ascii=False)
    assert client.delete(PREFIX + '/conversations/' + conversation['id']).status_code == 200
    with tenant_session(client.user_id) as db:
        assert list(db.scalars(select(Message).where(Message.conversation_id == conversation['id']))) == []


def test_invalid_batch_does_not_persist_any_attachments(chat):
    client, _, conversation = chat
    initial = client.get(PREFIX + '/conversations/' + conversation['id'] + '/messages').json()
    response = client.post(PREFIX + '/conversations/' + conversation['id'] + '/messages', json={'attachments': [file('good.md', b'task'), file('bad.json', b'{bad')], 'idempotency_key': str(uuid.uuid4())})
    assert response.status_code == 422
    data = client.get(PREFIX + '/conversations/' + conversation['id'] + '/messages').json()
    assert data == initial and data['runs'] == []
    injected = {**file('good.md', b'task'), 'text': 'client-forged-extraction'}
    assert client.post(PREFIX + '/conversations/' + conversation['id'] + '/messages', json={'attachments': [injected], 'idempotency_key': str(uuid.uuid4())}).status_code == 422


def test_forgotten_source_immediately_removes_attachment_content(chat):
    client, _, conversation = chat
    response = client.post(PREFIX + '/conversations/' + conversation['id'] + '/messages', json={'content': '我习惯早晨工作', 'attachments': [file('private.md', b'sensitive evidence')], 'idempotency_key': str(uuid.uuid4())})
    assert response.status_code == 202
    with tenant_session(client.user_id) as db:
        message = db.scalar(select(Message).where(Message.conversation_id == conversation['id'], Message.role == 'user'))
        memory = Memory(user_id=client.user_id, content='我习惯早晨工作', key='时段', source_message_id=message.id)
        db.add(memory); db.flush()
        identifier, message_id = memory.id, message.id
    assert client.post(PREFIX + '/memories/' + identifier + '/forget').status_code == 200
    with tenant_session(client.user_id) as db:
        assert db.get(Message, message_id).attachments == []


def test_verified_vision_sends_images_and_text_mode_never_fabricates_picture(chat, monkeypatch):
    import app.agent as agent
    client, conn, conversation = chat
    # Use a plain photograph-like color swatch: OCR returns no text.
    image = Image.new('RGB', (128, 128), 'blue')
    output = io.BytesIO(); image.save(output, format='PNG')
    body = {'content': '描述图片', 'attachments': [file('photo.png', output.getvalue())], 'idempotency_key': str(uuid.uuid4())}
    captured = []
    def respond(c, payload):
        captured.append(payload['messages'][-1]['content'])
        yield from output_stream({'answer': '图片说明'})
    monkeypatch.setattr(agent, 'stream_chat', respond)
    first = client.post(PREFIX + '/conversations/' + conversation['id'] + '/messages', json=body).json()
    execute_agent(client.user_id, first['id'])
    assert isinstance(captured[-1], str) and '不能猜测画面' in captured[-1]
    with tenant_session(client.user_id) as db:
        db.get(ProviderConnection, conn['id']).capabilities = {'test-model': {'chat': True, 'tools': True, 'vision': True}}
    second = client.post(PREFIX + '/conversations/' + conversation['id'] + '/messages', json={**body, 'idempotency_key': str(uuid.uuid4())}).json()
    execute_agent(client.user_id, second['id'])
    assert isinstance(captured[-1], list) and captured[-1][-1]['type'] == 'image_url'


def test_vision_probe_uses_actual_image_without_disclosing_color_answer(chat, monkeypatch):
    import app.providers as providers
    client, conn, _ = chat
    def probe(c, payload, user_id):
        assert isinstance(payload['messages'][0]['content'], list)
        parts = payload['messages'][0]['content']
        raw = base64.b64decode(parts[1]['image_url']['url'].split(',', 1)[1])
        with Image.open(io.BytesIO(raw)) as image:
            names = {(224, 0, 0): '红', (0, 176, 0): '绿', (0, 0, 224): '蓝'}
            answer = ','.join(names[image.getpixel((x, 32))] for x in (32, 96, 160))
        yield {'choices': [{'delta': {'content': answer}}]}
    monkeypatch.setattr(providers, 'tracked_probe', probe)
    response = client.post(PREFIX + '/provider-connections/' + conn['id'] + '/test-vision', json={'model': 'test-model'})
    assert response.status_code == 200 and response.json()['vision']
    with tenant_session(client.user_id) as db:
        cap = db.get(ProviderConnection, conn['id']).capabilities['test-model']
        assert cap['vision'] and cap['tools'] and cap['chat']
